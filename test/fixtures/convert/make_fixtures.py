#!/usr/bin/env python3
"""
Regenerate the small sample files used by test/test_convert_actions.py.

    .venv/bin/python test/fixtures/convert/make_fixtures.py [--audio]

Writes into this folder: sample.pptx, sample.xlsx, sample.xlsm, sample.epub and
(with --audio, needs kokoro) sample_dialogue.mp3. The cached formula values of
the .xlsx come from a LibreOffice round-trip; the tests derive .ppt / .xls
variants themselves (LibreOffice) so they are not stored here.
"""
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))


def _lo_convert(src: str, fmt: str, outdir: str) -> str:
    subprocess.run(['libreoffice', '--headless', '--convert-to', fmt, '--outdir', outdir, src],
                   check=True, capture_output=True, timeout=300)
    return os.path.join(outdir, os.path.splitext(os.path.basename(src))[0] + '.' + fmt.split(':')[0])


def make_pptx():
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    from PIL import Image

    prs = Presentation()
    # 1 — title slide + notes
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = 'Corporate Strategy'
    s.placeholders[1].text = 'Session 1 — Competitive advantage'
    s.notes_slide.notes_text_frame.text = 'Welcome everyone.\nStart with the Porter quote.'
    # 2 — bullets with levels
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = 'Five Forces'
    tf = s.placeholders[1].text_frame
    tf.text = 'Threat of new entrants'
    p = tf.add_paragraph(); p.text = 'Economies of scale'; p.level = 1
    p = tf.add_paragraph(); p.text = 'Bargaining power of buyers'
    s.notes_slide.notes_text_frame.text = 'Ask the class for an industry example.'
    # 3 — table
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = 'Market shares'
    t = s.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    for r, row in enumerate([['Firm', 'Share'], ['Alpha', '40%'], ['Beta | Gamma', '25%']]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    # 4 — chart
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = 'Revenue'
    cd = CategoryChartData(); cd.categories = ['2024', '2025']; cd.add_series('EUR m', (12.5, 14.0))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(6), Inches(4), cd)
    # 5 — image only
    s = prs.slides.add_slide(prs.slide_layouts[6])
    with tempfile.TemporaryDirectory() as tmp:
        img = os.path.join(tmp, 'px.png')
        Image.new('RGB', (40, 30), (200, 30, 30)).save(img)
        s.shapes.add_picture(img, Inches(1), Inches(1))
        # 6 — hidden slide
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = 'Backup slide'
        s.placeholders[1].text = 'Only shown on request'
        s._element.set('show', '0')
        prs.save(os.path.join(HERE, 'sample.pptx'))


def make_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = 'P&L'
    ws['A1'] = 'Item'; ws['B1'] = '2024'; ws['C1'] = '2025'
    ws['A2'] = 'Revenue'; ws['B2'] = 100; ws['C2'] = 120.5
    ws['A3'] = 'Costs'; ws['B3'] = 60; ws['C3'] = 70
    ws['A4'] = 'Margin'; ws['B4'] = '=B2-B3'; ws['C4'] = '=C2-C3'
    ws['A6'] = 'Growth | YoY'; ws['C6'] = '=C2/B2-1'
    ws.merge_cells('A8:C8'); ws['A8'] = 'Source: case study'
    ws2 = wb.create_sheet('Assumptions')
    ws2['B3'] = 'WACC'; ws2['C3'] = 0.08
    ws3 = wb.create_sheet('Scratch'); ws3['A1'] = 'hidden note'; ws3.sheet_state = 'hidden'
    wb.create_sheet('Empty')
    with tempfile.TemporaryDirectory() as tmp:
        raw = os.path.join(tmp, 'sample.xlsx')
        wb.save(raw)
        # round-trip through LibreOffice so formula results are cached in the file
        lo = os.path.join(tmp, 'lo'); os.mkdir(lo)
        shutil.copy(_lo_convert(raw, 'xlsx:Calc MS Excel 2007 XML', lo), os.path.join(HERE, 'sample.xlsx'))
        shutil.copy(os.path.join(HERE, 'sample.xlsx'), os.path.join(HERE, 'sample.xlsm'))


def make_epub():
    md = """---
title: Sample Strategy Book
lang: en
---

# Introduction

Strategy is about choices.

# Competitive Advantage

## Cost leadership

Be the cheapest producer.

## Differentiation

Be unique in ways buyers value.

# Conclusion

Choose one and commit.
"""
    with tempfile.TemporaryDirectory() as tmp:
        src = os.path.join(tmp, 'book.md')
        with open(src, 'w') as f:
            f.write(md)
        subprocess.run(['pandoc', src, '-o', os.path.join(HERE, 'sample.epub'),
                        '--split-level=1'], check=True)


def make_audio():
    import numpy as np
    import soundfile as sf
    from kokoro import KPipeline
    lines = [('af_heart', 'Good morning everyone. Today we discuss competitive advantage.'),
             ('am_michael', 'Thank you. Could you explain the difference between cost leadership and differentiation?'),
             ('af_heart', 'Of course. Cost leadership means being the cheapest producer in the industry.')]
    pipe = KPipeline(lang_code='a')
    sr, chunks = 24000, []
    for voice, text in lines:
        for _, _, audio in pipe(text, voice=voice):
            chunks.append(np.asarray(audio))
        chunks.append(np.zeros(int(sr * 0.8), dtype=np.float32))
    with tempfile.TemporaryDirectory() as tmp:
        wav = os.path.join(tmp, 'd.wav')
        sf.write(wav, np.concatenate(chunks), sr)
        subprocess.run(['ffmpeg', '-y', '-loglevel', 'error', '-i', wav, '-ac', '1', '-b:a', '48k',
                        os.path.join(HERE, 'sample_dialogue.mp3')], check=True)


if __name__ == '__main__':
    make_pptx()
    make_xlsx()
    make_epub()
    if '--audio' in sys.argv:
        make_audio()
    print('fixtures written to', HERE)
