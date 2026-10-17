#!/usr/bin/env python3
"""
ppt_to_md.py — Convert a PowerPoint deck (.pptx / .ppt) to Markdown.

Designed for knowledge extraction where exact locations matter:
- one `<!-- slide N -->` marker per slide, in deck order (N is 1-based)
- `## Slide N — <title>` heading (title placeholder, or "(untitled)")
- body text in reading order (top→bottom, left→right), bullet levels kept
- tables as GitHub pipe tables, charts as a small data table
- speaker notes in a delimited block:
      <!-- notes slide N -->
      **Speaker notes:**
      …
      <!-- /notes -->
- slides with pictures but no text are flagged `[image-only slide]`
- hidden slides are flagged `(hidden)` in the heading

.ppt files are converted to .pptx with LibreOffice first (temp dir, the input
folder is not touched).

Usage:
  python ppt_to_md.py <file.pptx|file.ppt> [-o OUT.md]
Default output: <basename>.md next to the input.
"""

import argparse
import os
import shutil
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER


SEP = '─' * 60
R   = '\033[0m'
B   = '\033[1m'
GRN = '\033[32m'
RED = '\033[31m'
DIM = '\033[2m'
HI  = '\033[36m'

_TITLE_PH = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
# Placeholders that carry no content of their own (rendered from the master).
_SKIP_PH  = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _clean(text: str) -> str:
    # \x0b = vertical tab python-pptx uses for soft line breaks
    return ' '.join(text.replace('\x0b', ' ').split())


def _cell(text: str) -> str:
    return _clean(text).replace('|', '\\|')


def _md_table(rows: list) -> list:
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [''] * (width - len(r)) for r in rows]
    out = ['| ' + ' | '.join(rows[0]) + ' |',
           '|' + '---|' * width]
    out += ['| ' + ' | '.join(r) + ' |' for r in rows[1:]]
    return out


def _ph_type(shape):
    if not shape.is_placeholder:
        return None
    try:
        return shape.placeholder_format.type
    except Exception:
        return None


_A_NS = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


class _RawShape:
    """Fallback for a shape element python-pptx cannot wrap (unregistered XML,
    e.g. some ink / content parts): keeps its position and raw `a:t` text."""
    shape_type = None
    is_placeholder = False
    has_text_frame = False

    def __init__(self, elm):
        off = next(elm.iter(f'{_A_NS}off'), None)
        self.top = int(off.get('y', 0)) if off is not None else 0
        self.left = int(off.get('x', 0)) if off is not None else 0
        paras = [''.join(t.text or '' for t in p.iter(f'{_A_NS}t')) for p in elm.iter(f'{_A_NS}p')]
        self.raw_lines = ['- ' + _clean(t) for t in paras if _clean(t)]


def _iter_shapes(shapes):
    try:
        elms = list(shapes._spTree.iter_shape_elms())
    except Exception:
        yield from shapes
        return
    for elm in elms:
        try:
            yield shapes._shape_factory(elm)
        except Exception:
            yield _RawShape(elm)


def _sorted_shapes(shapes):
    def key(s):
        return (s.top if s.top is not None else 0, s.left if s.left is not None else 0)
    return sorted(_iter_shapes(shapes), key=key)


def _text_frame_lines(tf) -> list:
    lines = []
    for p in tf.paragraphs:
        txt = _clean(''.join(r.text for r in p.runs) or p.text)
        if not txt:
            continue
        lines.append('  ' * (p.level or 0) + '- ' + txt)
    return lines


def _chart_lines(chart) -> list:
    out = []
    try:
        title = _clean(chart.chart_title.text_frame.text) if chart.has_title else ''
    except Exception:
        title = ''
    out.append(f'*Chart{": " + title if title else ""}*')
    try:
        plot = chart.plots[0]
        cats = [str(c) for c in plot.categories]
        series = list(plot.series)
        if series:
            rows = [[''] + [_cell(s.name or f'Series {i+1}') for i, s in enumerate(series)]]
            n = max(len(cats), max(len(list(s.values)) for s in series))
            for i in range(n):
                row = [_cell(cats[i]) if i < len(cats) else '']
                for s in series:
                    vals = list(s.values)
                    v = vals[i] if i < len(vals) else ''
                    row.append('' if v is None else _cell(f'{v:g}' if isinstance(v, float) else str(v)))
                rows.append(row)
            out += [''] + _md_table(rows)
    except Exception:
        pass
    return out


def _shape_blocks(shape, stats: dict) -> list:
    """Return a list of Markdown blocks (each a list of lines) for one shape."""
    if isinstance(shape, _RawShape):
        if shape.raw_lines:
            stats['text'] = True
            return [shape.raw_lines]
        return []
    try:
        st = shape.shape_type
    except Exception:
        st = None
    if st == MSO_SHAPE_TYPE.GROUP:
        blocks = []
        for sub in _sorted_shapes(shape.shapes):
            blocks += _shape_blocks(sub, stats)
        return blocks
    if st == MSO_SHAPE_TYPE.PICTURE or (shape.is_placeholder and hasattr(shape, 'image')):
        stats['pictures'] += 1
        return []
    if getattr(shape, 'has_table', False) and shape.has_table:
        rows = [[_cell(c.text) for c in r.cells] for r in shape.table.rows]
        stats['text'] = True
        return [_md_table(rows)]
    if getattr(shape, 'has_chart', False) and shape.has_chart:
        stats['text'] = True
        return [_chart_lines(shape.chart)]
    if shape.has_text_frame:
        lines = _text_frame_lines(shape.text_frame)
        if lines:
            stats['text'] = True
            return [lines]
    return []


def slide_to_md(n: int, slide) -> list:
    title = ''
    blocks = []
    stats = {'pictures': 0, 'text': False}
    for shape in _sorted_shapes(slide.shapes):
        pht = _ph_type(shape)
        if pht in _SKIP_PH:
            continue
        if pht in _TITLE_PH and not title and shape.has_text_frame:
            title = _clean(shape.text_frame.text)
            if title:
                continue
        blocks += _shape_blocks(shape, stats)

    hidden = slide._element.get('show') == '0'
    out = [f'<!-- slide {n} -->',
           f'## Slide {n} — {title or "(untitled)"}{" (hidden)" if hidden else ""}', '']
    for b in blocks:
        out += b + ['']
    if not title and not stats['text']:
        out += ['[image-only slide]' if stats['pictures'] else '[empty slide]', '']
    elif stats['pictures']:
        out += [f'*[{stats["pictures"]} image(s) on slide]*', '']

    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame
        paras = [_clean(p.text) for p in notes.paragraphs] if notes is not None else []
        paras = [p for p in paras if p]
        if paras:
            out += [f'<!-- notes slide {n} -->', '**Speaker notes:**', '']
            out += ['\n\n'.join(paras), '', '<!-- /notes -->', '']
    return out


def _ppt_to_pptx(ppt: str, tmpdir: str) -> str:
    cmd = ['libreoffice', '--headless', '--convert-to', 'pptx', '--outdir', tmpdir, ppt]
    subprocess.run(cmd, check=True, stdout=sys.stderr, stderr=sys.stderr, timeout=600)
    out = os.path.join(tmpdir, os.path.splitext(os.path.basename(ppt))[0] + '.pptx')
    if not os.path.isfile(out):
        raise RuntimeError('LibreOffice did not produce a .pptx')
    return out


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description='PowerPoint → Markdown (slides, tables, notes)')
    ap.add_argument('input')
    ap.add_argument('-o', '--output', help='output .md (default: <base>.md beside input)')
    args = ap.parse_args()

    src = os.path.abspath(args.input)
    if not os.path.isfile(src):
        print(f'{RED}File not found:{R} {src}', file=sys.stderr)
        sys.exit(1)
    out_md = os.path.abspath(args.output) if args.output else os.path.splitext(src)[0] + '.md'

    print(SEP)
    print(f'  {B}PPT → Markdown{R}')
    print(SEP)
    print(f'  {DIM}Input :{R}  {HI}{os.path.basename(src)}{R}')
    print(f'  {DIM}Output:{R}  {HI}{out_md}{R}')

    tmpdir = tempfile.mkdtemp(prefix='ppt_to_md_')
    try:
        path = src
        if src.lower().endswith('.ppt'):
            print(f'  {DIM}Converting .ppt → .pptx (LibreOffice)…{R}')
            path = _ppt_to_pptx(src, tmpdir)
        prs = Presentation(path)
        lines = []
        for n, slide in enumerate(prs.slides, 1):
            lines += slide_to_md(n, slide)
    except Exception as e:
        print(f'  {RED}Error:{R} {e}', file=sys.stderr)
        sys.exit(1)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

    title = os.path.splitext(os.path.basename(src))[0]
    header = [f'# {title}', '', f'<!-- source: {os.path.basename(src)} | slides: {n if prs.slides else 0} -->', '']
    with open(out_md, 'w', encoding='utf-8') as f:
        f.write('\n'.join(header + lines).rstrip() + '\n')

    print(f'  {GRN}✓{R} {len(prs.slides)} slides → {os.path.basename(out_md)}')
    print(SEP)


if __name__ == '__main__':
    main()
