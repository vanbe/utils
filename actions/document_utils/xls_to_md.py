#!/usr/bin/env python3
"""
xls_to_md.py — Convert a spreadsheet (.xlsx / .xlsm / .xls) to Markdown.

Designed for knowledge extraction where exact cell locations matter:
- one `<!-- sheet N: <name> -->` marker + `## Sheet N — <name>` per sheet
  (workbook order; hidden sheets flagged `(hidden)`)
- the used range as a pipe table whose header row is the column letters and
  whose first column is the row number → every value is citable as e.g. `B7`
- cell values are the cached results Excel/LibreOffice saved with the file
- a `### Formulas` list per sheet: `` `B7`: `=SUM(B2:B6)` → 123 ``

.xls files are converted to .xlsx with LibreOffice first (temp dir, the input
folder is not touched).

Usage:
  python xls_to_md.py <file.xlsx|.xlsm|.xls> [-o OUT.md] [--max-rows N] [--max-formulas N]
Default output: <basename>.md next to the input.
"""

import argparse
import datetime as dt
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl
from openpyxl.utils import get_column_letter


SEP = '─' * 60
R   = '\033[0m'
B   = '\033[1m'
GRN = '\033[32m'
RED = '\033[31m'
DIM = '\033[2m'
HI  = '\033[36m'


def _fmt(v) -> str:
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, float):
        return f'{v:.10g}'
    if isinstance(v, dt.datetime):
        return v.isoformat(sep=' ', timespec='minutes') if (v.hour or v.minute) else v.date().isoformat()
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    return ' '.join(str(v).split())


def _cell(v) -> str:
    return _fmt(v).replace('|', '\\|')


def _used_bounds(ws_f, ws_v):
    """Tight (min_row, min_col, max_row, max_col) over non-empty cells, or None."""
    rows, cols = [], []
    for ws in (ws_f, ws_v):
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and str(c.value).strip() != '':
                    rows.append(c.row)
                    cols.append(c.column)
    if not rows:
        return None
    return min(rows), min(cols), max(rows), max(cols)


def sheet_to_md(n: int, ws_f, ws_v, max_rows: int, max_formulas: int) -> list:
    hidden = ws_f.sheet_state != 'visible'
    out = [f'<!-- sheet {n}: {ws_f.title} -->',
           f'## Sheet {n} — {ws_f.title}{" (hidden)" if hidden else ""}', '']

    bounds = _used_bounds(ws_f, ws_v)
    if bounds is None:
        return out + ['[empty sheet]', '']
    r0, c0, r1, c1 = bounds
    ref = f'{get_column_letter(c0)}{r0}:{get_column_letter(c1)}{r1}'
    out += [f'Used range: `{ref}`', '']

    letters = [get_column_letter(c) for c in range(c0, c1 + 1)]
    out.append('| | ' + ' | '.join(letters) + ' |')
    out.append('|---|' + '---|' * len(letters))
    formulas = []
    shown = 0
    for r in range(r0, r1 + 1):
        vals = []
        for c in range(c0, c1 + 1):
            f = ws_f.cell(row=r, column=c).value
            v = ws_v.cell(row=r, column=c).value
            if isinstance(f, str) and f.startswith('='):
                formulas.append((f'{get_column_letter(c)}{r}', f, v))
            elif v is None and f is not None:
                v = f        # array formula objects etc. → show the raw value
            vals.append(_cell(v))
        if not any(vals):
            continue
        if shown < max_rows:
            out.append(f'| {r} | ' + ' | '.join(vals) + ' |')
        shown += 1
    if shown > max_rows:
        out += ['', f'*[{shown - max_rows} more non-empty rows truncated (--max-rows {max_rows})]*']
    out.append('')

    if ws_f.merged_cells.ranges:
        out += ['Merged cells: ' + ', '.join(f'`{m}`' for m in ws_f.merged_cells.ranges), '']

    if formulas:
        out += ['### Formulas', '']
        for addr, f, v in formulas[:max_formulas]:
            res = _fmt(v)
            out.append(f'- `{addr}`: `{f}`' + (f' → {res}' if res else ''))
        if len(formulas) > max_formulas:
            out.append(f'- *[{len(formulas) - max_formulas} more formulas truncated (--max-formulas {max_formulas})]*')
        out.append('')
    return out


def _xls_to_xlsx(xls: str, tmpdir: str) -> str:
    cmd = ['libreoffice', '--headless', '--convert-to', 'xlsx', '--outdir', tmpdir, xls]
    subprocess.run(cmd, check=True, stdout=sys.stderr, stderr=sys.stderr, timeout=600)
    out = os.path.join(tmpdir, os.path.splitext(os.path.basename(xls))[0] + '.xlsx')
    if not os.path.isfile(out):
        raise RuntimeError('LibreOffice did not produce a .xlsx')
    return out


def main():
    ap = argparse.ArgumentParser(description='Spreadsheet → Markdown (tables + formulas)')
    ap.add_argument('input')
    ap.add_argument('-o', '--output', help='output .md (default: <base>.md beside input)')
    ap.add_argument('--max-rows', type=int, default=2000, help='non-empty rows per sheet (default 2000)')
    ap.add_argument('--max-formulas', type=int, default=2000, help='formulas listed per sheet (default 2000)')
    args = ap.parse_args()

    src = os.path.abspath(args.input)
    if not os.path.isfile(src):
        print(f'{RED}File not found:{R} {src}', file=sys.stderr)
        sys.exit(1)
    out_md = os.path.abspath(args.output) if args.output else os.path.splitext(src)[0] + '.md'

    print(SEP)
    print(f'  {B}XLS → Markdown{R}')
    print(SEP)
    print(f'  {DIM}Input :{R}  {HI}{os.path.basename(src)}{R}')
    print(f'  {DIM}Output:{R}  {HI}{out_md}{R}')

    tmpdir = tempfile.mkdtemp(prefix='xls_to_md_')
    try:
        path = src
        if src.lower().endswith('.xls'):
            print(f'  {DIM}Converting .xls → .xlsx (LibreOffice)…{R}')
            path = _xls_to_xlsx(src, tmpdir)
        wb_f = openpyxl.load_workbook(path, data_only=False)
        wb_v = openpyxl.load_workbook(path, data_only=True)
        lines = []
        for n, name in enumerate(wb_f.sheetnames, 1):
            ws_f = wb_f[name]
            if not hasattr(ws_f, 'iter_rows'):      # chartsheet
                lines += [f'<!-- sheet {n}: {name} -->', f'## Sheet {n} — {name}', '', '[chart sheet]', '']
                continue
            lines += sheet_to_md(n, ws_f, wb_v[name], args.max_rows, args.max_formulas)
    except Exception as e:
        print(f'  {RED}Error:{R} {e}', file=sys.stderr)
        sys.exit(1)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)

    title = os.path.splitext(os.path.basename(src))[0]
    header = [f'# {title}', '', f'<!-- source: {os.path.basename(src)} | sheets: {len(wb_f.sheetnames)} -->', '']
    with open(out_md, 'w', encoding='utf-8') as f:
        f.write('\n'.join(header + lines).rstrip() + '\n')

    print(f'  {GRN}✓{R} {len(wb_f.sheetnames)} sheets → {os.path.basename(out_md)}')
    print(SEP)


if __name__ == '__main__':
    main()
